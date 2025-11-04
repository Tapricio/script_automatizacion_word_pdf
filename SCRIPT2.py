import os
import re
import pandas as pd
from datetime import datetime
from tkinter import Tk, filedialog
from docx import Document
import win32com.client
import comtypes.client

# Obtener la fecha actual en formato DD-MM-YYYY
fecha_hoy = datetime.today().strftime('%d-%m-%Y')
idTemplate = 11271

# Carpeta base
base_folder_name = rf".\output\cartas reconocer"

# Oculta la ventana principal de Tkinter
root = Tk()
root.withdraw()

# Selección de archivo Excel
file_path = filedialog.askopenfilename(
    title="Selecciona un archivo Excel",
    filetypes=[("Excel files", "*.xlsx *.xls")]
)

# Listas de errores
datosErroneos = []
id = 11271

# Función para limpiar nombres de archivo
def limpiar_nombre_archivo(nombre):
    return re.sub(r'[\\:*?"<>|]', '', nombre)

# Función para reemplazar texto por regex
def docx_replace_regex(doc_obj, regex, replace):
    for p in doc_obj.paragraphs:
        if regex.search(p.text):
            inline = p.runs
            for i in range(len(inline)):
                if regex.search(inline[i].text):
                    inline[i].text = regex.sub(replace, inline[i].text)
    for table in doc_obj.tables:
        for row in table.rows:
            for cell in row.cells:
                docx_replace_regex(cell, regex, replace)

# Leer Excel y procesar filas
if file_path:
    df = pd.read_excel(file_path, header=1)
    for index, row in df.iterrows():
        if pd.notna(df.iloc[index, 2]):  # Validar que hay paciente
            
            # Nombre
            try:
                nombre = re.sub(r'\s+', ' ', df.iloc[index, 1].strip())
            except Exception:
                datosErroneos.append(f"index: {index}, data - Nombre: {df.iloc[index, 1]} Rut: {df.iloc[index, 2]}")
                continue

            # RUT
            try:
                rut_response = df.iloc[index, 2].replace(" ", "").replace(".", "")
                rut_base = rut_response[:-2]
                verificador = rut_response[-1]
                rut = "{:,}".format(int(rut_base)).replace(",", ".") + "-" + verificador
                rut = rut.upper()
            except Exception:
                datosErroneos.append(f"index: {index}, data - Nombre: {df.iloc[index, 1]} Rut: {df.iloc[index, 2]}")
                continue

            print(f"Nombre: {nombre} Rut: {rut} ID: {id}")

            # Abrir plantilla Word
            doc = Document("CartaTemplate.docx")

            # Reemplazos base (sin ValorTemplate)
            docx_replace_regex(doc, re.compile(r"NombreTemplate"), nombre)
            docx_replace_regex(doc, re.compile(r"RutTemplate"), rut)
            docx_replace_regex(doc, re.compile(r"IdTemplate"), str(id))
            docx_replace_regex(doc, re.compile(r"FechaEmisionTemplate"), fecha_hoy)

            # Guardar Word
            nombre_archivo_seguro = limpiar_nombre_archivo(nombre)
            word_folder = r".\output\cartas reconocer\word"
            os.makedirs(word_folder, exist_ok=True)
            ruta_guardado = os.path.join(word_folder, f"{nombre_archivo_seguro} {rut}.docx")
            doc.save(ruta_guardado)

            id += 1

# Mostrar errores
if datosErroneos:
    print("------------------------")
    print("ERROR:")
    for error in datosErroneos:
        print(error)
    print("------------------------")
